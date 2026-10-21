import logging
from datetime import date, timedelta
from decimal import Decimal

from django.contrib import messages
from django.contrib.messages.views import SuccessMessageMixin
from django.db.models import Case, Count, DateField, ExpressionWrapper, F, IntegerField, Q, Value, When
from django.db.models.functions import Coalesce
from django.http import JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.utils.decorators import method_decorator
from django.views import View
from django.views.decorators.http import require_POST
from django.views.generic import ListView, DetailView, TemplateView
from django.views.generic.edit import UpdateView

from farm.forms import TreatmentForm, TreatmentProductFormSet
from farm.mixins import BaseSecureViewMixin
from farm.models import Field, Product, ProductType, Treatment, TreatmentProduct
from farm.services import save_treatment_with_products, get_shopping_list, clone_treatment

logger = logging.getLogger(__name__)


class TreatmentListView(BaseSecureViewMixin, ListView):
    model = Treatment
    template_name = 'farm/treatments/treatment_list.html'
    context_object_name = 'treatments'
    paginate_by = 24
    ordering = ['-date']
    # Estados de la pestaña «Por hacer»
    TODO_STATUSES = [Treatment.STATUS_PENDING, Treatment.STATUS_DELAYED]

    # ── Etiquetas legibles para el selector de orden ──────────────────────────
    SORT_LABELS = {
        'urgency':   'Por urgencia',
        'date_desc': 'Fecha: más reciente primero',
        'date_asc':  'Fecha: más antigua primero',
        'name_asc':  'Nombre A-Z',
        'name_desc': 'Nombre Z-A',
    }

    def _get_filter_dates(self):
        """Devuelve (date_from, date_to, from_is_default, to_is_default).

        Si el usuario no ha enviado los parámetros, se usa el año en curso
        como rango por defecto para que el listado esté acotado.
        """
        current_year = date.today().year
        from_in_get = 'date_from' in self.request.GET
        to_in_get = 'date_to' in self.request.GET
        date_from = self.request.GET.get('date_from') if from_in_get else f'{current_year}-01-01'
        date_to = self.request.GET.get('date_to') if to_in_get else f'{current_year}-12-31'
        return date_from, date_to, not from_in_get, not to_in_get

    def _filtered_queryset(self, apply_status=True):
        """Queryset con todos los filtros del GET aplicados (sin anotar ni ordenar).

        Con ``apply_status=False`` se ignora el filtro de estado: sirve para contar
        los pendientes de la pestaña «Por hacer» aunque se esté viendo «Completados».
        """
        queryset = (
            super()
            .get_queryset()
            .select_related("field", "machine")
            .prefetch_related("treatmentproduct_set__product__product_type")
        )

        field_ids = self.request.GET.getlist('field')
        type_filters = [v for v in self.request.GET.getlist('type') if v]
        status_filters = [v for v in self.request.GET.getlist('status') if v]
        product_ids = self.request.GET.getlist('products')
        date_from, date_to, _, _ = self._get_filter_dates()
        product_type_filters = self.request.GET.getlist('product_types')
        search_query = self.request.GET.get('q', '').strip()

        if field_ids:
            queryset = queryset.filter(field__id__in=field_ids)
        if type_filters:
            queryset = queryset.filter(type__in=type_filters)
        if product_ids:
            queryset = queryset.filter(products__id__in=product_ids)
        # Rango de fechas: los atrasados (sin finalizar y con fecha pasada) aparecen
        # siempre, aunque caigan fuera del rango, porque son lo más urgente.
        today = date.today()
        date_range = Q()
        if date_from:
            date_range &= Q(date__gte=date_from)
        if date_to:
            date_range &= Q(date__lte=date_to)
        if date_range:
            queryset = queryset.filter(date_range | Q(finish_date__isnull=True, date__lt=today))
        if product_type_filters:
            queryset = queryset.filter(products__product_type__in=product_type_filters)
        if apply_status:
            if status_filters:
                queryset = queryset.filter(status__in=status_filters)
            elif 'status' not in self.request.GET:
                # Sin parámetro status → pestaña «Por hacer» (pendientes + atrasados)
                queryset = queryset.filter(status__in=self.TODO_STATUSES)
            # status presente pero vacío → todos los estados
        if search_query:
            queryset = queryset.filter(name__icontains=search_query)
        return queryset

    def get_queryset(self):
        queryset = self._filtered_queryset()
        today = date.today()
        sort = self.request.GET.get('sort', 'urgency')

        sort_map = {
            'date_desc': '-effective_date',
            'date_asc':  'effective_date',
            'name_asc':  'name',
            'name_desc': '-name',
        }
        queryset = queryset.annotate(
            effective_date=ExpressionWrapper(
                Coalesce('finish_date', 'date'),
                output_field=DateField()
            ),
            # Grupo de urgencia calculado a partir de fechas (no del campo status,
            # que puede estar desfasado hasta que se guarde el tratamiento).
            urgency_group=Case(
                When(finish_date__isnull=False, then=Value(3)),            # Completados
                When(date__lt=today, then=Value(0)),                       # Atrasados
                When(date__lte=today + timedelta(days=7), then=Value(1)),  # Próximos 7 días
                default=Value(2),                                          # Más adelante
                output_field=IntegerField(),
            ),
            pending_date=Case(When(finish_date__isnull=True, then=F('date')), output_field=DateField()),
        )
        if sort == 'urgency':
            # Atrasados de más antiguo a más reciente; completados de más reciente a más antiguo
            return queryset.order_by(
                'urgency_group',
                F('pending_date').asc(nulls_last=True),
                F('finish_date').desc(nulls_last=True),
                'name',
            ).distinct()
        order_field = sort_map.get(sort, '-effective_date')
        return queryset.order_by(order_field).distinct()

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)

        selected_fields = self.request.GET.getlist('field')
        selected_types = [v for v in self.request.GET.getlist('type') if v]
        selected_statuses = [v for v in self.request.GET.getlist('status') if v]
        selected_products = self.request.GET.getlist('products')
        selected_product_types = self.request.GET.getlist('product_types')
        date_from, date_to, date_from_is_default, date_to_is_default = self._get_filter_dates()
        search_query = self.request.GET.get('q', '').strip()
        sort = self.request.GET.get('sort', 'urgency')

        context['fields'] = Field.ownership_objects.get_queryset_for_user(self.request.user)
        context['products'] = Product.ownership_objects.get_queryset_for_user(self.request.user)
        context['type_choices'] = Treatment.TYPE_CHOICES
        context['status_choices'] = Treatment.STATUS_CHOICES
        context['selected_fields'] = selected_fields
        context['selected_types'] = selected_types
        context['selected_statuses'] = selected_statuses
        context['selected_products'] = selected_products
        context['date_from'] = date_from
        context['date_to'] = date_to
        context['date_from_is_default'] = date_from_is_default
        context['date_to_is_default'] = date_to_is_default
        context['product_types'] = ProductType.ownership_objects.get_queryset_for_user(self.request.user)
        context['selected_product_types'] = selected_product_types
        context['total_count'] = self.object_list.count()
        context['show_groups'] = (sort == 'urgency')
        # Conteo por grupo sobre el total filtrado (no sobre la página), en una sola consulta
        group_counts = {0: 0, 1: 0, 2: 0, 3: 0}
        for row in self.object_list.order_by().values('urgency_group').annotate(n=Count('id', distinct=True)):
            group_counts[row['urgency_group']] = row['n']
        context['group_counts'] = group_counts
        context['available_fields'] = Field.ownership_objects.get_queryset_for_user(self.request.user)
        context['search_query'] = search_query
        context['sort'] = sort
        context['sort_label'] = self.SORT_LABELS.get(sort, self.SORT_LABELS['urgency'])

        # ── Pestañas de estado: «Por hacer» (pendientes + atrasados) / «Completados» ──
        # Se apoyan en el parámetro `status`, así siguen valiendo los enlaces que ya
        # existen desde otras pantallas (p. ej. ?status=delayed).
        status_set = set(selected_statuses)
        if status_set == {Treatment.STATUS_COMPLETED}:
            status_tab = 'done'
        elif not status_set or status_set <= set(self.TODO_STATUSES):
            status_tab = 'todo'
        else:
            status_tab = 'all'  # combinación antigua (p. ej. todos los estados)
        context['status_tab'] = status_tab
        # Filtro de estado más fino que la pestaña (solo pendientes o solo atrasados)
        context['status_subfilter'] = (
            dict(Treatment.STATUS_CHOICES).get(selected_statuses[0], "")
            if status_tab == 'todo' and len(status_set) == 1 else ''
        )
        context['todo_count'] = (
            self._filtered_queryset(apply_status=False)
            .filter(status__in=self.TODO_STATUSES).order_by().distinct().count()
        )
        base_params = self.request.GET.copy()
        for key in ('status', 'page'):
            base_params.pop(key, None)
        context['todo_url'] = '?' + base_params.urlencode()
        done_params = base_params.copy()
        done_params.setlist('status', [Treatment.STATUS_COMPLETED])
        context['done_url'] = '?' + done_params.urlencode()

        active_filter_count = (
            len(selected_fields)
            + len(selected_types)
            + len(selected_products)
            + len(selected_product_types)
            + (1 if date_from or date_to else 0)
            + (1 if search_query else 0)
        )
        context['active_filter_count'] = active_filter_count
        # Filtros del panel (sin búsqueda ni fechas por defecto): para el contador del botón «Filtros» en móvil
        context['panel_filter_count'] = (
            len(selected_fields)
            + len(selected_types)
            + len(selected_products)
            + len(selected_product_types)
            + (0 if date_from_is_default and date_to_is_default else 1)
        )

        query_params = self.request.GET.copy()
        if 'page' in query_params:
            del query_params['page']
        context['query_params'] = query_params.urlencode()
        query_params.pop('sort', None)
        context['sort_query'] = query_params.urlencode()

        return context


class TreatmentDetailView(BaseSecureViewMixin, DetailView):
    model = Treatment
    template_name = 'farm/treatments/treatment_detail.html'
    context_object_name = 'treatment'

    def get_queryset(self):
        return (
            super()
            .get_queryset()
            .select_related("field", "machine")
            .prefetch_related("treatmentproduct_set__product__product_type")
        )

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['status_choices'] = Treatment.STATUS_CHOICES

        products = self.object.treatmentproduct_set.all()
        context['products'] = products

        total_cost = sum(product.total_price for product in products)
        cost_per_ha = 0
        if self.object.effective_area and self.object.effective_area > 0:
            cost_per_ha = total_cost / Decimal(self.object.effective_area)

        context['total_cost'] = total_cost
        context['cost_per_ha'] = cost_per_ha
        context['has_prices'] = any(p.unit_price > 0 for p in products)

        # Cubas (solo pulverización con máquina y mojado; fertirrigación no lleva)
        context['load_info'] = self.object.calculate_machine_loads()

        # Tono de la pastilla de fecha: mismo criterio que los grupos del listado
        # 0 atrasado · 1 próximos 7 días · 2 más adelante · 3 completado
        today = date.today()
        if self.object.finish_date:
            tone = 3
        elif self.object.date < today:
            tone = 0
        elif self.object.date <= today + timedelta(days=7):
            tone = 1
        else:
            tone = 2
        context['tone'] = tone

        context['available_fields'] = (
            Field.ownership_objects
            .get_queryset_for_user(self.request.user)
            .exclude(pk=self.object.field_id)
        )

        return context


class TreatmentFormView(BaseSecureViewMixin, SuccessMessageMixin, UpdateView):
    model = Treatment
    form_class = TreatmentForm
    template_name = 'farm/treatments/treatment_form.html'

    def get_object(self, queryset=None):
        if 'pk' not in self.kwargs:
            return None
        return super().get_object(queryset)

    def get_initial(self):
        initial = super().get_initial()
        # ?date=YYYY-MM-DD (p. ej. desde el calendario) preselecciona la fecha al crear
        preset = self.request.GET.get('date')
        if preset and 'pk' not in self.kwargs:
            try:
                initial['date'] = date.fromisoformat(preset).isoformat()
            except ValueError:
                pass
        return initial

    def get_success_url(self):
        return reverse('treatment-detail', kwargs={'pk': self.object.id})

    def get_success_message(self, cleaned_data):
        return f"Tratamiento '{cleaned_data['name']}' guardado exitosamente"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['is_edit'] = bool(self.object and self.object.pk)
        context['form_title'] = 'Editar tratamiento' if context['is_edit'] else 'Nuevo tratamiento'

        # Usar el formset ya validado (con errores) si existe; si no, crear uno nuevo.
        if hasattr(self, '_products_formset'):
            context['products_formset'] = self._products_formset
        elif self.request.POST:
            context['products_formset'] = TreatmentProductFormSet(self.request.POST, instance=self.object)
        else:
            context['products_formset'] = TreatmentProductFormSet(instance=self.object)

        # Indicar al template si hay errores en el formset para navegar al paso 2
        fs = context['products_formset']
        context['has_formset_errors'] = bool(
            any(fs.errors) or fs.non_form_errors()
        ) if hasattr(self, '_products_formset') else False

        return context

    def post(self, request, *args, **kwargs):
        self.object = self.get_object()
        form = self.get_form()
        # Guardar referencia al formset validado para reutilizarlo en get_context_data
        self._products_formset = TreatmentProductFormSet(request.POST, instance=self.object)
        if form.is_valid() and self._products_formset.is_valid():
            self.object = save_treatment_with_products(form, self._products_formset)
            return self.form_valid(form)
        return self.form_invalid(form)


class TreatmentCalendarView(BaseSecureViewMixin, TemplateView):
    template_name = 'farm/treatments/calendar.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['fields'] = Field.ownership_objects.get_queryset_for_user(self.request.user)
        context['treatment_types'] = Treatment.TYPE_CHOICES
        return context


@method_decorator(require_POST, name='dispatch')
class FinishTreatmentView(BaseSecureViewMixin, View):
    def post(self, request, pk):
        treatment = get_object_or_404(Treatment, pk=pk)
        finish_date = request.POST.get('finish_date')
        real_water_used = request.POST.get('real_water_used')

        if not finish_date:
            return JsonResponse({'success': False}, status=400)

        from datetime import date
        try:
            finish_date_obj = date.fromisoformat(finish_date) if finish_date else None
        except (ValueError, TypeError):
            finish_date_obj = None

        treatment.finish_treatment(finish_date_obj, real_water_used)

        from accounts.models import Notification
        from accounts.notification_service import notify_org_users
        from farm.services import build_treatment_email_html
        products = list(treatment.treatmentproduct_set.select_related('product').order_by('position'))
        product_names = ', '.join(tp.product.name for tp in products) if products else 'Sin productos'
        finish_date_display = finish_date_obj.strftime('%d/%m/%Y') if finish_date_obj else ''
        extra_rows = [('Fecha de finalización', finish_date_display)]
        if treatment.is_spraying() and treatment.real_water_per_ha:
            extra_rows.append(('Mojado real', f'{treatment.real_water_per_ha} L/ha'))
        if treatment.is_partial_treatment:
            extra_rows.append(('Área tratada', f'{treatment.applied_area} ha de {treatment.field.area} ha'))
        if treatment.zone_notes:
            extra_rows.append(('Zona', treatment.zone_notes))
        notify_org_users(
            event_type=Notification.EVENT_TREATMENT_FINISHED,
            title=f'Tratamiento finalizado: {treatment.name}',
            body=(
                f'El tratamiento "{treatment.name}" en {treatment.field.name} ha sido marcado como completado.'
                + (f' Productos: {product_names}.' if products else '')
            ),
            link=reverse('treatment-detail', kwargs={'pk': treatment.pk}),
            organization=treatment.organization,
            html_body=build_treatment_email_html(treatment, extra_rows=extra_rows),
        )

        return JsonResponse({'success': True})


@method_decorator(require_POST, name='dispatch')
class DeleteTreatmentView(BaseSecureViewMixin, View):
    def post(self, request, pk):
        treatment = get_object_or_404(Treatment, pk=pk)
        treatment.delete()
        messages.success(request, f'Tratamiento "{treatment.name}" eliminado')
        return redirect('treatment-list')


class CloneTreatmentView(BaseSecureViewMixin, View):
    template_name = 'farm/treatments/treatment_clone.html'

    def get_treatment(self, pk):
        return get_object_or_404(Treatment, pk=pk)

    def get(self, request, pk):
        treatment = self.get_treatment(pk)
        available_fields = (
            Field.ownership_objects
            .get_queryset_for_user(request.user)
            .exclude(pk=treatment.field_id)
        )
        return render(request, self.template_name, {
            'treatment': treatment,
            'available_fields': available_fields,
        })

    def post(self, request, pk):
        treatment = self.get_treatment(pk)

        field_ids = request.POST.getlist('field_id')
        date_str = request.POST.get('date')

        if not field_ids:
            messages.error(request, 'Debes seleccionar al menos una parcela de destino.')
            return redirect('treatment-clone', pk=pk)

        cloned = []
        for field_id in field_ids:
            target_field = get_object_or_404(Field, pk=field_id, organization=request.user.organization)
            new_name = request.POST.get(f'name_{field_id}', '').strip() or None
            per_field_date = request.POST.get(f'date_{field_id}', '').strip()
            try:
                new_date = date.fromisoformat(per_field_date) if per_field_date else (
                    date.fromisoformat(date_str) if date_str else None
                )
            except ValueError:
                new_date = None
            cloned.append(clone_treatment(treatment, target_field, new_date, new_name))

        if len(cloned) == 1:
            messages.success(request, f'Tratamiento clonado en {cloned[0].field.name}.')
            return redirect('treatment-detail', pk=cloned[0].pk)
        else:
            field_names = ', '.join(t.field.name for t in cloned)
            messages.success(request, f'Tratamiento clonado en: {field_names}.')
            return redirect('treatment-list')


class RepeatTreatmentView(BaseSecureViewMixin, View):
    """
    Genera copias de un tratamiento en la misma parcela para un conjunto de fechas.
    Permite seleccionar fechas individualmente o mediante un patrón recurrente
    (ej: todos los viernes entre dos fechas).
    """
    template_name = 'farm/treatments/treatment_repeat.html'

    def get_treatment(self, pk):
        return get_object_or_404(Treatment, pk=pk, organization=self.request.user.organization)

    def get(self, request, pk):
        treatment = self.get_treatment(pk)
        return render(request, self.template_name, {'treatment': treatment})

    def post(self, request, pk):
        treatment = self.get_treatment(pk)

        # La lista de fechas llega como valores múltiples de 'dates'
        raw_dates = request.POST.getlist('dates')

        if not raw_dates:
            messages.error(request, 'Debes seleccionar al menos una fecha.')
            return render(request, self.template_name, {'treatment': treatment})

        created = []
        errors = []
        for raw in raw_dates:
            raw = raw.strip()
            if not raw:
                continue
            try:
                new_date = date.fromisoformat(raw)
            except ValueError:
                errors.append(raw)
                continue
            created.append(treatment.clone_to_field(treatment.field, new_date=new_date))

        if errors:
            messages.warning(request, f'Algunas fechas no eran válidas y se ignoraron: {", ".join(errors)}')

        if not created:
            messages.error(request, 'No se pudo crear ningún tratamiento. Revisa las fechas.')
            return render(request, self.template_name, {'treatment': treatment})

        if len(created) == 1:
            messages.success(request, f'Tratamiento repetido el {created[0].date.strftime("%d/%m/%Y")}.')
            return redirect('treatment-detail', pk=created[0].pk)

        messages.success(
            request,
            f'Se han creado {len(created)} tratamientos para las fechas seleccionadas.'
        )
        return redirect('treatment-list')


class TreatmentExportView(BaseSecureViewMixin, DetailView):
    model = Treatment
    template_name = 'farm/treatments/treatment_export.html'
    context_object_name = 'treatment'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['products'] = TreatmentProduct.ownership_objects.get_queryset_for_user(self.request.user).filter(
            treatment=self.object)
        context['now'] = timezone.now()
        return context


class ShoppingListView(BaseSecureViewMixin, ListView):
    template_name = 'farm/treatments/shopping_list.html'
    context_object_name = 'product_items'

    def _selected_treatments(self):
        return [tid for tid in self.request.GET.getlist('treatment') if tid.isdigit()]

    def _selected_fields(self):
        return [fid for fid in self.request.GET.getlist('field') if fid.isdigit()]

    def get_queryset(self):
        return get_shopping_list(
            self.request.user,
            field_ids=self._selected_fields() or None,
            treatment_ids=self._selected_treatments() or None,
        )

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['fields'] = Field.objects.filter(organization=self.request.user.organization)
        context['selected_fields'] = self._selected_fields()
        context['selected_treatments'] = self._selected_treatments()
        context['available_treatments'] = (
            Treatment.ownership_objects
            .get_queryset_for_user(self.request.user)
            .filter(status__in=['pending', 'delayed'])
            .select_related('field')
            .order_by('date')
        )
        total_price = sum(item['total_price'] for item in self.object_list)
        context['total_price'] = round(total_price, 2)
        context['total_count'] = len(self.object_list)
        return context


class TreatmentExcelExportView(BaseSecureViewMixin, View):
    """Exporta tratamientos filtrados a Excel para contabilidad."""
    
    def get(self, request):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from django.http import HttpResponse
        
        # Reutilizar la lógica de filtrado de TreatmentListView
        list_view = TreatmentListView()
        list_view.request = request
        queryset = list_view.get_queryset()
        
        # Crear workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Tratamientos"
        
        # Estilos
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Headers
        headers = [
            "Fecha",
            "Parcela",
            "Tipo",
            "Nombre",
            "Estado",
            "Productos",
            "Costo Total",
            "Área (ha)",
            "Notas"
        ]
        ws.append(headers)
        
        # Aplicar estilos al header
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
        
        # Datos
        for treatment in queryset:
            # Calcular costo total de productos
            total_cost = sum(tp.total_price for tp in treatment.treatmentproduct_set.all())
            
            # Obtener productos
            products = ", ".join([
                f"{tp.product.name}" for tp in treatment.treatmentproduct_set.all()
            ])
            
            row = [
                treatment.date.strftime("%d/%m/%Y") if treatment.date else "",
                treatment.field.name,
                dict(Treatment.TYPE_CHOICES).get(treatment.type, treatment.type),
                treatment.name,
                dict(Treatment.STATUS_CHOICES).get(treatment.status, treatment.status),
                products,
                f"{total_cost:.2f}€",
                treatment.applied_area or "",
                treatment.zone_notes or "",
            ]
            ws.append(row)
            
            # Aplicar estilos a datos
            for cell in ws[ws.max_row]:
                cell.border = border
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        
        # Ajustar ancho de columnas
        ws.column_dimensions['A'].width = 12  # Fecha
        ws.column_dimensions['B'].width = 15  # Parcela
        ws.column_dimensions['C'].width = 12  # Tipo
        ws.column_dimensions['D'].width = 20  # Nombre
        ws.column_dimensions['E'].width = 12  # Estado
        ws.column_dimensions['F'].width = 30  # Productos
        ws.column_dimensions['G'].width = 12  # Costo
        ws.column_dimensions['H'].width = 10  # Área
        ws.column_dimensions['I'].width = 25  # Notas
        
        # Crear respuesta HTTP
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="tratamientos.xlsx"'
        wb.save(response)
        return response


